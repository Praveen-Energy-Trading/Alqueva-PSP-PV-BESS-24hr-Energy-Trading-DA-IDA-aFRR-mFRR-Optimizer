"""
test_excel_layout.py — generated Excel sheets are laid out by their content: narrow columns for small numbers under long headers, wrapped
headers, two-decimal floats, label/value sheets sized by their longest cell, and no value changed.
"""
from __future__ import annotations

import datetime as dt

from openpyxl import Workbook

from common_layer.utilities import excel_layout as xl


def table(title_rows: int = 0):
    wb = Workbook()
    ws = wb.active
    for i in range(title_rows):
        ws.append([f"A long title line {i} that spans the whole sheet and must not widen column A", None, None])
    ws.append(["Hour", "Reservoir_inflow_m3h", "Market description"])
    for h in range(1, 25):
        ws.append([h, 340853.4 + h / 7.0, "Day-Ahead auction (OMIE)" if h == 1 else "x"])
    return wb, ws


def test_the_width_follows_the_data_not_the_header():
    wb, ws = table()
    xl.fit_columns(ws, header_row=1)
    assert ws.column_dimensions["A"].width <= 10                              # two-digit hours under the header 'Hour'
    assert 10 <= ws.column_dimensions["B"].width <= 14                        # 340,853.54 under a 20-character header: not 22 wide
    assert ws.column_dimensions["C"].width >= len("Day-Ahead auction (OMIE)") + 2   # the longest value decides
    assert ws["B1"].alignment.wrap_text and ws.row_dimensions[1].height > 15  # the long header wraps onto more lines


def test_rows_above_the_header_do_not_set_the_widths_and_no_value_changes():
    wb, ws = table(title_rows=2)
    before = [[c.value for c in row] for row in ws.iter_rows()]
    xl.fit_columns(ws, header_row=3)
    assert ws.column_dimensions["A"].width <= 10
    assert [[c.value for c in row] for row in ws.iter_rows()] == before


def test_floats_get_two_decimals_but_other_formats_are_kept():
    wb, ws = table()
    ws["B3"].number_format = "0.000"
    xl.fit_columns(ws, header_row=1)
    assert ws["B2"].number_format == xl.FLOAT_FORMAT and ws["B3"].number_format == "0.000" and ws["A2"].number_format == "General"


def test_cells_are_centred_and_long_text_stays_left():
    wb, ws = table()
    ws["C3"] = "A long description of more than forty characters that is not centred"
    xl.fit_columns(ws, header_row=1)
    assert ws["A2"].alignment.horizontal == "center" and ws["B2"].alignment.horizontal == "center" and ws["B1"].alignment.horizontal == "center"
    assert ws["C3"].alignment.horizontal == "left" and ws["C3"].alignment.wrap_text
    kv = Workbook().active
    kv.append(["Label", 12.5])
    xl.fit_key_value_sheet(kv)
    assert kv["B1"].alignment.horizontal == "center" and kv["A1"].alignment.horizontal != "center"


def test_shown_length_counts_what_is_displayed():
    assert xl.shown_length(1234567.891) == len("1,234,567.89") and xl.shown_length(0.2083333333333333) == 4
    assert xl.shown_length(True) == 4 and xl.shown_length(None) == 0 and xl.shown_length(dt.date(2026, 1, 1)) == 10
    assert xl.shown_length("two\nlines of text") == len("lines of text")


def test_a_label_value_sheet_is_sized_by_its_longest_cell_and_a_table_workbook_is_frozen():
    wb = Workbook()
    kv = wb.active
    kv.title = "Summary"
    kv.append(["Backtest summary with a very long title in A1 that spills to the right"])
    kv.append(["Days", 365])
    kv.append(["Average realised revenue, real OMIE price (EUR)", 123456.789])
    t = wb.create_sheet("Table")
    t.append(["date", "feasible", "objective_eur"])
    for i in range(5):
        t.append([f"2026-10-0{i + 1}", True, 1000.5 + i])
    xl.format_workbook(wb)
    assert kv.column_dimensions["A"].width == len("Average realised revenue, real OMIE price (EUR)") + 2     # the A1 title does not widen the column
    assert t.freeze_panes == "A2" and t["A1"].alignment.wrap_text and t.column_dimensions["B"].width <= 12
